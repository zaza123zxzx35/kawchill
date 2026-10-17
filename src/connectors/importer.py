from __future__ import annotations

import csv
import io
from datetime import datetime

from openpyxl import load_workbook

from src.labels import AFFILIATION, EXAM_GENERAL, EXAM_SPECIAL, SUBJECT
from src.parsing import clean_text, org_level, parse_number

HEADER_MAP = {
    "exam_year": "exam_year",
    "ปีสอบ": "exam_year",
    "exam_type": "exam_type",
    "ประเภทการสอบ": "exam_type",
    "subject": "subject",
    "วิชาเอก": "subject",
    "affiliation": "affiliation",
    "สังกัด": "affiliation",
    "org_level": "org_level",
    "ระดับเขต": "org_level",
    "area_name": "area_name",
    "เขตพื้นที่": "area_name",
    "openings": "openings",
    "เปิดรับ": "openings",
    "applicants": "applicants",
    "ผู้สมัคร": "applicants",
    "listed_count": "listed_count",
    "ขึ้นบัญชี": "listed_count",
    "appointed_count": "appointed_count",
    "เรียกบรรจุ": "appointed_count",
    "source_url": "source_url",
    "url": "source_url",
    "reference_date": "reference_date",
    "วันที่อ้างอิง": "reference_date",
    "notes": "notes",
    "หมายเหตุ": "notes",
}

EXAM_ALIASES = {
    "general": EXAM_GENERAL,
    "การคัดเลือกทั่วไป": EXAM_GENERAL,
    "สอบแข่งขันทั่วไป": EXAM_GENERAL,
    "special": EXAM_SPECIAL,
    "กรณีพิเศษ": EXAM_SPECIAL,
}


def _rows_from_csv(payload: bytes) -> list[dict]:
    text = payload.decode("utf-8-sig", "replace")
    return list(csv.DictReader(io.StringIO(text)))


def _rows_from_xlsx(payload: bytes) -> list[dict]:
    workbook = load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
    sheet = workbook.active
    grid = list(sheet.iter_rows(values_only=True))
    if not grid:
        return []
    headers = [clean_text(str(cell)) if cell is not None else "" for cell in grid[0]]
    rows = []
    for raw in grid[1:]:
        rows.append({headers[index]: raw[index] if index < len(raw) else None for index in range(len(headers))})
    return rows


def parse_import_file(filename: str, payload: bytes) -> tuple[list[dict], list[str]]:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        raw_rows = _rows_from_xlsx(payload)
    elif name.endswith(".csv"):
        raw_rows = _rows_from_csv(payload)
    else:
        return [], ["รองรับเฉพาะไฟล์ CSV หรือ XLSX"]
    records: list[dict] = []
    warnings: list[str] = []
    seen: set[tuple] = set()
    for index, raw in enumerate(raw_rows, start=2):
        mapped: dict[str, str] = {}
        for key, value in raw.items():
            target = HEADER_MAP.get(clean_text(str(key)))
            if target:
                mapped[target] = "" if value is None else clean_text(str(value))
        if not any(mapped.values()):
            continue
        area_name = mapped.get("area_name", "")
        year_number, year_known, year_warn = parse_number(mapped.get("exam_year"))
        exam_type = EXAM_ALIASES.get(mapped.get("exam_type", "").strip())
        subject = mapped.get("subject") or SUBJECT
        if not area_name or not year_known or exam_type is None:
            warnings.append(f"แถว {index}: ต้องมีปีสอบ ประเภทการสอบ และเขตพื้นที่")
            continue
        if year_warn:
            warnings.append(f"แถว {index}: {year_warn}")
            continue
        openings, openings_known, openings_warn = parse_number(mapped.get("openings"))
        applicants, applicants_known, applicants_warn = parse_number(mapped.get("applicants"))
        listed, listed_known, listed_warn = parse_number(mapped.get("listed_count"))
        appointed, appointed_known, appointed_warn = parse_number(mapped.get("appointed_count"))
        for warning in (openings_warn, applicants_warn, listed_warn, appointed_warn):
            if warning:
                warnings.append(f"แถว {index}: {warning}")
        remaining = None
        if listed_known and appointed_known:
            if appointed > listed:
                warnings.append(f"แถว {index}: เรียกบรรจุมากกว่าจำนวนขึ้นบัญชี จึงไม่คำนวณคงเหลือ")
            else:
                remaining = listed - appointed
        key = (exam_type, subject, year_number, area_name)
        if key in seen:
            warnings.append(f"แถว {index}: ข้อมูลซ้ำกับแถวก่อนหน้า")
            continue
        seen.add(key)
        reference = mapped.get("reference_date") or None
        records.append(
            {
                "area_name": area_name,
                "area_key": f"import:{area_name}",
                "org_level": mapped.get("org_level") or org_level(area_name),
                "exam_year": year_number,
                "exam_type": exam_type,
                "subject": subject,
                "affiliation": mapped.get("affiliation") or AFFILIATION,
                "opened": True if openings_known and openings > 0 else (False if openings_known and openings == 0 else None),
                "openings": openings if openings_known else None,
                "applicants": applicants if applicants_known else None,
                "applicants_known": applicants_known,
                "listed_count": listed if listed_known else None,
                "appointed_count": appointed if appointed_known else None,
                "remaining_count": remaining,
                "detail_url": None,
                "source_url": mapped.get("source_url") or "local:csv",
                "reference_at": reference,
                "notes": mapped.get("notes") or "",
                "imported_at": datetime.now().astimezone().isoformat(timespec="seconds"),
            }
        )
    if not records and not warnings:
        warnings.append("ไม่พบแถวข้อมูลในไฟล์")
    return records, warnings
